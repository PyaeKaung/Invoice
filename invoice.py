import tkinter as tk
from tkinter import Canvas, StringVar, Listbox, messagebox
from tkcalendar import Calendar
from PIL import Image, ImageTk
from openpyxl import load_workbook
from datetime import datetime, timedelta

class OptionMenu:
    def __init__(self, master, variable, options):
        self.master = master
        self.variable = variable
        self.options = options

        self.button = tk.Button(master, text=variable.get(), command=self.show_options, bg='black', fg='white', font=("Helvetica", 20), width=24)
        self.button.place()

        self.listbox = Listbox(master, bg='black', fg='white', font=("Helvetica", 20), height=min(3, len(options)))
        for option in options:
            self.listbox.insert(tk.END, option)
        self.listbox.bind("<<ListboxSelect>>", self.select_option)
        self.listbox.place_forget()

    def show_options(self):
        if self.listbox.winfo_ismapped():
            self.listbox.place_forget()
        else:
            x = self.button.winfo_x()
            y = self.button.winfo_y() + self.button.winfo_height()
            self.listbox.place(x=x, y=y)
            self.listbox.lift()

    def select_option(self, event):
        selected_option = self.listbox.get(self.listbox.curselection())
        self.variable.set(selected_option)
        self.button.config(text=selected_option)
        self.listbox.place_forget()

def submit_data(name, cfm, room_quantity, rows, room_type, extras, date_option, calendar, remarks):
    try:
        workbook = load_workbook('template.xlsx')
        sheet = workbook.active
        
        sheet['D7'] = name
        sheet['D8'] = cfm
        row_number = int(rows)
        target_row = 11 + row_number - 1
        
        sheet[f'C{target_row}'] = room_quantity
        sheet[f'D{target_row}'] = room_type
        sheet[f'E{target_row}'] = extras

        selected_checkin_date = calendar.get_date()  
        checkin_date = datetime.strptime(selected_checkin_date, '%m/%d/%y')

        selected_checkout_date = calendar.get_date()  
        checkout_date = datetime.strptime(selected_checkout_date, '%m/%d/%y')

        selected_invoice_date = calendar.get_date()  
        invoice_date = datetime.strptime(selected_invoice_date, '%m/%d/%y')

        if date_option == "Check-in Date":
            sheet[f'F{target_row}'] = checkin_date

        elif date_option == "Check-out Date":
            sheet[f'G{target_row}'] = checkout_date

        elif date_option == "Invoice Date":
            sheet['M7'] = invoice_date

        if date_option == "Check-in Date":
            sheet[f'F{target_row}'].number_format = 'DD/MM/YYYY'
        elif date_option == "Check-out Date":
            sheet[f'G{target_row}'].number_format = 'DD/MM/YYYY'
        elif date_option == "Invoice Date":
            sheet['M7'].number_format = 'DD/MM/YYYY'

        sheet['F29'].number_format = 'DD/MM/YYYY'
        
        sheet['C18'] = remarks
        sheet['G29'].number_format = 'DD/MM/YYYY'
        sheet['G28'].number_format = 'DD/MM/YYYY'

        starting_value = {
            "-": 0,
            "A-Suite": 235000,
            "DRV": 140000,
            "DGV": 100000,
            "SUP": 85000
        }.get(room_type, 0)
        sheet[f'I{target_row}'] = starting_value
        extra_value = {
            "-": 0,
            "Extra bed": 40000,
            "Extra breakfast": 12000 
        }.get(extras, 0)
        sheet[f'J{target_row}'] = extra_value
        
        workbook.save('template.xlsx')
        
        messagebox.showinfo("Success", "Data submitted successfully!")
    except ValueError as ve:
        messagebox.showerror("Error", f"ValueError: {ve}")
    except Exception as e:
        messagebox.showerror("Error", f"An unexpected error occurred: {e}")

def main():
    root = tk.Tk()
    root.title("Invoice Application")
    root.geometry("800x1200")
    root.resizable(False, False)
    canvas = Canvas(root, width=800, height=1200)
    canvas.pack()

    background_image = Image.open("invoicebackground.jpg")
    background_image = background_image.resize((800, 1200), Image.NEAREST)
    background_photo = ImageTk.PhotoImage(background_image)
    canvas.create_image(0, 0, anchor=tk.NW, image=background_photo)

    text_frame1 = tk.Frame(root, width=400, height=40, bg="white", bd=2, relief="solid")
    text_frame1.place(x=170, y=40)

    text_field1 = tk.Entry(text_frame1, font=("Helvetica", 16), bd=0)
    text_field1.pack(fill=tk.BOTH, expand=True)

    label1 = tk.Label(root, text="Name:", font=("Helvetica", 16), bg='black', fg='white')
    label1.place(x=60, y=40)

    text_frame2 = tk.Frame(root, width=400, height=40, bg="white", bd=2, relief="solid")
    text_frame2.place(x=170, y=90)

    text_field2 = tk.Entry(text_frame2, font=("Helvetica", 16), bd=0)
    text_field2.pack(fill=tk.BOTH, expand=True)

    label2 = tk.Label(root, text="CFM #", font=("Helvetica", 16), bg='black', fg='white')
    label2.place(x=60, y=90)

    option3_var = StringVar(root)
    option3_var.set("Room quantity")
    options3 = [str(i) for i in range(1, 100)]
    option_menu3 = OptionMenu(root, option3_var, options3)
    option_menu3.button.place(x=60, y=150)

    option1_var = StringVar(root)
    option1_var.set("Room type")
    options1 = ["-", "A-Suite", "DRV", "DGV", "SUP"]
    option_menu1 = OptionMenu(root, option1_var, options1)
    option_menu1.button.place(x=60, y=230)

    option2_var = StringVar(root)
    option2_var.set("Extras")
    options2 = ["-", "Extra bed", "Extra breakfast"]
    option_menu2 = OptionMenu(root, option2_var, options2)
    option_menu2.button.place(x=60, y=310)

    option4_var = StringVar(root)
    option4_var.set("Row")
    options4 = [str(i) for i in range(1, 7)]
    option_menu4 = OptionMenu(root, option4_var, options4)
    option_menu4.button.place(x=60, y=390)

    date_options = ["Check-in Date", "Check-out Date", "Invoice Date"]
    date_var = StringVar(root)
    date_var.set(date_options[0])
    date_menu = OptionMenu(root, date_var, date_options)
    date_menu.button.place(x=420, y=80)

    calendar = Calendar(root, selectmode='day', year=datetime.now().year, month=datetime.now().month, day=datetime.now().day)
    calendar.place(x=480, y=150, width=300, height=300)

    submit_button = tk.Button(root, text="Submit", command=lambda: (
        submit_data(
            text_field1.get(), 
            text_field2.get(), 
            option3_var.get(),
            option4_var.get(), 
            option1_var.get(), 
            option2_var.get(),
            date_var.get(),
            calendar,
            remarks_text.get("1.0", tk.END).strip()
        )
    ), bg='green', fg='white', font=("Helvetica", 20))
    submit_button.place(x=680, y=550)

    remarks_label = tk.Label(root, text="Remarks:", font=("Helvetica", 16), bg='black', fg='white')
    remarks_label.place(x=60, y=470)
    remarks_text = tk.Text(root, font=("Helvetica", 16), bd=2, wrap=tk.WORD, height=8, width=50)
    remarks_text.place(x=60, y=520)

    root.mainloop()

if __name__ == "__main__":
    main()
